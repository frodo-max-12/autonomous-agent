"""
Import vendors from vendors_template.xlsx into the Vendor table.

Usage:
    python scripts/import_vendors.py                      # imports ./vendors_template.xlsx
    python scripts/import_vendors.py path/to/vendors.xlsx

Upserts by (name + first email) so re-running with an updated sheet refreshes rows
instead of duplicating them. Rows with no name+email are skipped.
"""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl
from config.settings import get_settings
from core.database import init_database, Vendor


def _split(val) -> list[str]:
    if val is None:
        return []
    return [p.strip() for p in str(val).replace(";", ",").split(",") if p.strip()]


def _truthy(val, default=True) -> bool:
    if val is None:
        return default
    return str(val).strip().lower() in ("yes", "y", "true", "1", "active")


def main(path: str):
    xlsx = Path(path)
    if not xlsx.is_absolute():
        xlsx = ROOT / xlsx
    if not xlsx.exists():
        print(f"[import_vendors] file not found: {xlsx}")
        return 1

    settings = get_settings()
    _, Session = init_database(settings.database_url)
    session = Session()

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Vendors"] if "Vendors" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]

    def col(row, name):
        return row[headers.index(name)] if name in headers else None

    added = updated = skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        name = (col(row, "vendor_name") or "").strip() if col(row, "vendor_name") else ""
        email = (col(row, "vendor_email") or "").strip() if col(row, "vendor_email") else ""
        if not name or not email:
            skipped += 1
            continue

        first_email = _split(email)[0] if _split(email) else email
        existing = (session.query(Vendor)
                    .filter(Vendor.name == name, Vendor.email.like(f"{first_email}%"))
                    .first())
        fields = dict(
            name=name,
            email=email,
            brands=[b.upper() for b in _split(col(row, "brands"))],
            categories=_split(col(row, "categories")),
            currency=(str(col(row, "currency")).strip().upper() if col(row, "currency") else "USD"),
            country=(str(col(row, "country")).strip() if col(row, "country") else None),
            priority=(int(col(row, "priority")) if col(row, "priority") else 3),
            active=_truthy(col(row, "active")),
            notes=(str(col(row, "notes")).strip() if col(row, "notes") else None),
        )
        if existing:
            for k, v in fields.items():
                setattr(existing, k, v)
            updated += 1
        else:
            session.add(Vendor(**fields))
            added += 1

    session.commit()
    total = session.query(Vendor).count()
    session.close()
    print(f"[import_vendors] added={added} updated={updated} skipped={skipped} | total vendors now={total}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "vendors_template.xlsx"))
